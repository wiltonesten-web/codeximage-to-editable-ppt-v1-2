#!/usr/bin/env python3
"""Build one refined slide from source-coordinate crops, text, and simple shapes."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


def color(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(*bytes.fromhex(value))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build one refined PPTX page from a JSON specification.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--reference-dir", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--pptx-out", type=Path, required=True)
    parser.add_argument("--page-id", required=True)
    parser.add_argument("--page-index", type=int, required=True)
    parser.add_argument("--source-input", type=Path, required=True)
    parser.add_argument("--spec-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    spec_path = args.spec_dir / f"{args.source_input.stem}.json"
    if not spec_path.exists():
        raise SystemExit(f"Missing page specification: {spec_path}")
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    args.outdir.mkdir(parents=True, exist_ok=True)
    assets = args.outdir / "split_png_elements_refined" / args.page_id
    assets.mkdir(parents=True, exist_ok=True)

    source = Image.open(args.source).convert("RGB")
    source_w, source_h = source.size
    if [source_w, source_h] != list(spec.get("source_size", [source_w, source_h])):
        raise SystemExit(f"Source size changed: expected {spec.get('source_size')}, got {[source_w, source_h]}")

    slide_w_in = float(spec.get("slide_width_in", 13.333333))
    source_ratio = source_w / source_h
    if abs(source_ratio - (16 / 9)) <= 0.02:
        slide_h_in = float(spec.get("slide_height_in", 7.5))
    else:
        slide_h_in = float(spec.get("slide_height_in", slide_w_in / source_ratio))
    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = color(spec.get("background", "#FFFFFF"))

    def emu_x(px: float) -> int:
        return int(round(px / source_w * prs.slide_width))

    def emu_y(px: float) -> int:
        return int(round(px / source_h * prs.slide_height))

    records: list[dict] = []
    allowed_text_assets: list[str] = []
    overlay = source.copy()
    draw = ImageDraw.Draw(overlay)
    palette = {"crop": "#00A67E", "text": "#D12C2C", "shape": "#7A42F4"}

    for index, item in enumerate(spec.get("elements", []), start=1):
        kind = item["type"]
        bbox = [float(v) for v in item["bbox_px"]]
        x, y, w, h = bbox
        if w <= 0 or h <= 0:
            raise SystemExit(f"Invalid bbox for element {index}: {bbox}")
        target = [float(v) for v in item.get("target_bbox_px", bbox)]
        tx, ty, tw, th = target
        element_id = item.get("id", f"{args.page_id}_{kind}_{index:03d}")
        asset_name = ""

        if kind == "crop":
            if w * h >= source_w * source_h * 0.95:
                raise SystemExit(f"Whole-slide crop is forbidden in refined mode: {element_id}")
            crop = source.crop((round(x), round(y), round(x + w), round(y + h)))
            asset_name = f"{element_id}.png"
            asset_path = assets / asset_name
            crop.save(asset_path)
            slide.shapes.add_picture(str(asset_path), emu_x(tx), emu_y(ty), width=emu_x(tw), height=emu_y(th))
            if bool(item.get("allow_embedded_text", False)):
                allowed_text_assets.append(asset_name)
        elif kind == "text":
            box = slide.shapes.add_textbox(emu_x(tx), emu_y(ty), emu_x(tw), emu_y(th))
            frame = box.text_frame
            frame.clear()
            frame.margin_left = frame.margin_right = Inches(float(item.get("margin_x_in", 0.02)))
            frame.margin_top = frame.margin_bottom = Inches(float(item.get("margin_y_in", 0.0)))
            frame.vertical_anchor = getattr(MSO_ANCHOR, item.get("vertical", "MIDDLE").upper())
            paragraph = frame.paragraphs[0]
            paragraph.alignment = getattr(PP_ALIGN, item.get("align", "CENTER").upper())
            run = paragraph.add_run(); run.text = item["text"]
            run.font.name = item.get("font", "Microsoft YaHei")
            run.font.size = Pt(float(item.get("font_size", 24)))
            run.font.bold = bool(item.get("bold", False))
            run.font.color.rgb = color(item.get("color", "#000000"))
        elif kind == "shape":
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu_x(tx), emu_y(ty), emu_x(tw), emu_y(th))
            shape.fill.solid(); shape.fill.fore_color.rgb = color(item.get("fill", "#FFFFFF"))
            if item.get("line"):
                shape.line.color.rgb = color(item["line"])
                shape.line.width = Pt(float(item.get("line_width", 1)))
            else:
                shape.line.fill.background()
        else:
            raise SystemExit(f"Unsupported element type: {kind}")

        draw.rectangle((round(x), round(y), round(x + w), round(y + h)), outline=palette[kind], width=max(2, source_w // 900))
        records.append({
            "page_id": args.page_id,
            "element_id": element_id,
            "element_type": item.get("element_type", kind),
            "file_name": asset_name,
            "is_text_only": "yes" if kind == "text" else "no",
            "allow_embedded_text": "yes" if bool(item.get("allow_embedded_text", False)) else "no",
            "source_bbox_px": [x, y, w, h],
            "pptx_target_bbox_px": [tx, ty, tw, th],
            "editable": kind in {"text", "shape"},
            "text": item.get("text", ""),
        })

    args.pptx_out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.pptx_out)
    (args.outdir / "visual_elements_manifest.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.outdir / "asset_text_allowlist.json").write_text(
        json.dumps({"allowed_files": allowed_text_assets}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (args.outdir / "visual_elements_manifest.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0].keys()))
        writer.writeheader(); writer.writerows(records)
    review = args.outdir / "review"
    review.mkdir(exist_ok=True)
    overlay.save(review / f"{args.page_id}_refined_overlay.png")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
