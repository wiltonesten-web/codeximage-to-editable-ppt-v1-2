from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


SCRIPT = Path(__file__).with_name("parallel_v1_pages.py")
SPEC = importlib.util.spec_from_file_location("parallel_v1_pages", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def make_page(path: Path, label: str, width_delta: int = 0) -> tuple[int, int, int, int]:
    deck = Presentation()
    deck.slide_width = int(deck.slide_width) + width_delta
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    shape.text = label
    geometry = (int(shape.left), int(shape.top), int(shape.width), int(shape.height))
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)
    return geometry


def make_picture_overlap_page(path: Path, asset: Path, full_slide: bool) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if full_slide:
        slide.shapes.add_picture(str(asset), 0, 0, deck.slide_width, deck.slide_height)
    else:
        slide.shapes.add_picture(str(asset), Inches(0.5), Inches(0.5), Inches(5), Inches(4))
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text = "ResidualText"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)


def make_terminal_report(path: Path, source: Path, status: str = "passed") -> None:
    render = path.parent / "powerpoint_render" / "slide_01.png"
    render.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (16, 9), "white").save(render)
    write_json(
        path,
        {
            "status": status,
            "renderer": "microsoft_powerpoint",
            "mode": "final_visual",
            "slides": [
                {
                    "slide": 1,
                    "status": status,
                    "renderer": "microsoft_powerpoint",
                    "rendered_path": str(render),
                    "source_path": str(source),
                    "mean_absolute_error": 1.25,
                    "changed_pixel_ratio": 0.01,
                    "text_overflow_count": 0,
                    "off_slide_shape_count": 0,
                    "issues": "",
                }
            ],
        },
    )


def make_branded_raster_page(path: Path, body_color: tuple[int, int, int]) -> None:
    image = Image.new("RGB", (160, 90), body_color)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 159, 10), fill=(82, 24, 128))
    draw.ellipse((132, 2, 139, 9), outline="white", width=1)
    draw.rectangle((143, 3, 155, 4), fill="white")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


class ParallelV1PagesTests(unittest.TestCase):
    def test_contract_source_hash_ignores_text_line_endings(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / "SKILL.md"
            contract = root / "page_worker_contract.md"
            source.write_bytes(b"line one\r\nline two\r\n")
            contract.write_text("contract", encoding="utf-8")
            expected = hashlib.sha256(b"line one\nline two\n").hexdigest()
            with patch.object(MODULE, "V1_CONTRACT_SOURCE_HASHES", {"SKILL.md": expected}):
                actual = MODULE.verify_worker_contract_sources(root, contract)
            self.assertEqual(expected, actual["SKILL.md"])

    def test_preflight_groups_reusable_branded_mastheads_once(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            rows = []
            for index, color in enumerate(((255, 255, 255), (245, 250, 255), (255, 248, 245)), start=1):
                source = root / f"slide_{index:02d}.png"
                make_branded_raster_page(source, color)
                rows.append(MODULE.analyze_preflight_input(source, index))
            plan = MODULE.plan_shared_masthead_reuse(rows)
            self.assertEqual("reuse_available", plan["status"])
            self.assertEqual(1, len(plan["groups"]))
            self.assertEqual(3, len(plan["groups"][0]["members"]))
            self.assertAlmostEqual(
                plan["groups"][0]["crop_fraction"],
                plan["groups"][0]["normalized_placement"][3],
            )
            self.assertGreater(plan["groups"][0]["crop_fraction"], 0.10)
            self.assertLess(plan["groups"][0]["crop_fraction"], 0.14)

    def test_shared_masthead_gate_checks_hash_position_size_and_crop(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "masthead.png"
            Image.new("RGB", (160, 11), (82, 24, 128)).save(asset)
            expected = {
                "group_id": "masthead_group_001",
                "asset_path": str(asset),
                "asset_sha256": MODULE.sha256_file(asset),
                "normalized_placement": [0.0, 0.0, 1.0, 11 / 90],
            }
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(
                str(asset), 0, 0, deck.slide_width, round(deck.slide_height * 11 / 90)
            )
            passed_path = root / "passed.pptx"
            deck.save(passed_path)
            issues, report = MODULE.shared_masthead_integrity_issues(passed_path, expected)
            self.assertEqual([], issues)
            self.assertEqual("passed", report["status"])

            shifted = Presentation()
            shifted_slide = shifted.slides.add_slide(shifted.slide_layouts[6])
            shifted_slide.shapes.add_picture(
                str(asset), 0, Inches(0.1), shifted.slide_width, round(shifted.slide_height * 11 / 90)
            )
            shifted_path = root / "shifted.pptx"
            shifted.save(shifted_path)
            issues, report = MODULE.shared_masthead_integrity_issues(shifted_path, expected)
            self.assertIn("shared_masthead_placement_or_crop_mismatch", issues)
            self.assertEqual("failed", report["status"])

            overlapping = Presentation()
            overlapping_slide = overlapping.slides.add_slide(overlapping.slide_layouts[6])
            overlapping_slide.shapes.add_picture(
                str(asset), 0, 0, overlapping.slide_width, round(overlapping.slide_height * 11 / 90)
            )
            title = overlapping_slide.shapes.add_textbox(
                Inches(0.5), Inches(0.4), Inches(5), Inches(0.8)
            )
            title.text = "Title colliding with masthead"
            overlapping_path = root / "overlapping.pptx"
            overlapping.save(overlapping_path)
            issues, report = MODULE.shared_masthead_integrity_issues(overlapping_path, expected)
            self.assertTrue(
                any(item.startswith("shared_masthead_overlaps_editable_text") for item in issues)
            )
            self.assertEqual("failed", report["status"])

    def test_prepare_materializes_one_shared_masthead_for_all_group_members(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            inputs = []
            for index, color in enumerate(((255, 255, 255), (245, 250, 255)), start=1):
                source = root / f"slide_{index:02d}.png"
                make_branded_raster_page(source, color)
                inputs.append(source)
            outdir = root / "output"
            report_path = root / "preflight.json"
            capacity = {
                "assessment_kind": "test",
                "logical_cpu_count": 8,
                "total_memory_gb": 16.0,
                "available_memory_gb": 12.0,
                "active_powerpoint_processes": 0,
                "assumptions": {},
                "capacity_by_resource": {"cpu": 2, "available_memory": 3, "powerpoint": 2},
                "hardware_safe_worker_limit": 2,
            }
            with patch.object(MODULE, "detect_system_capacity", return_value=capacity):
                MODULE.command_preflight(
                    argparse.Namespace(
                        inputs=inputs,
                        report=report_path,
                        intended_outdir=outdir,
                        available_worker_slots=2,
                        root_task_count=1,
                        cpu_threads_per_worker=4,
                        memory_gb_per_worker=4.0,
                        powerpoint_worker_cap=2,
                    )
                )
            MODULE.command_prepare(
                argparse.Namespace(
                    inputs=inputs,
                    outdir=outdir,
                    workers=2,
                    python=Path(sys.executable),
                    v1_skill=MODULE.DEFAULT_V1_SKILL,
                    worker_contract=MODULE.DEFAULT_WORKER_CONTRACT,
                    extractor=MODULE.DEFAULT_EXTRACTOR,
                    validator=MODULE.DEFAULT_VALIDATOR,
                    dpi=300,
                    timeout=60,
                    preflight_report=report_path,
                    confirmed=True,
                    capacity_override_confirmed=False,
                    clean=True,
                )
            )
            manifest = json.loads((outdir / "parallel_v1_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(1, len(manifest["shared_masthead_groups"]))
            assignments = [task["shared_masthead"] for task in manifest["tasks"]]
            self.assertEqual(2, len(assignments))
            self.assertEqual(1, len({item["asset_path"] for item in assignments}))
            self.assertEqual(1, len({item["asset_sha256"] for item in assignments}))
            for task in manifest["tasks"]:
                request = Path(task["worker_request"]).read_text(encoding="utf-8")
                self.assertIn("--brand-masthead-status shared_faithful --shared-masthead-used", request)

    def test_system_capacity_uses_cpu_memory_and_powerpoint_minimum(self) -> None:
        gib = 1024**3
        with (
            patch.object(MODULE.os, "cpu_count", return_value=16),
            patch.object(MODULE, "system_memory_bytes", return_value=(32 * gib, 20 * gib)),
            patch.object(MODULE, "active_powerpoint_process_count", return_value=1),
        ):
            capacity = MODULE.detect_system_capacity(
                cpu_threads_per_worker=4,
                memory_gb_per_worker=4.0,
                powerpoint_worker_cap=4,
            )
        self.assertEqual(4, capacity["capacity_by_resource"]["cpu"])
        self.assertEqual(5, capacity["capacity_by_resource"]["available_memory"])
        self.assertEqual(3, capacity["capacity_by_resource"]["powerpoint"])
        self.assertEqual(3, capacity["hardware_safe_worker_limit"])

    def test_six_worker_timing_gate_requires_observed_peak_six(self) -> None:
        metrics = {"peak_concurrent_workers": 3, "measured_speedup": 2.0}
        failed = MODULE.timing_acceptance(metrics, page_count=6, worker_limit=6, min_speedup=1.4)
        self.assertEqual(6, failed["required_peak_concurrent_workers"])
        self.assertEqual("failed", failed["status"])
        metrics["peak_concurrent_workers"] = 6
        passed = MODULE.timing_acceptance(metrics, page_count=6, worker_limit=6, min_speedup=1.4)
        self.assertEqual("passed", passed["status"])

    def test_terminal_report_rejects_failed_page(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / "source.png"
            Image.new("RGB", (16, 9), "white").save(source)
            report = root / "terminal" / "layout_quality_report.json"
            make_terminal_report(report, source, status="failed")
            issues, _ = MODULE.terminal_page_validation_issues(report, source)
            self.assertIn("terminal_validation_not_passed", issues)
            self.assertIn("terminal_validation_slide_not_passed", issues)

    def test_record_uses_inline_review_evidence_without_review_bundle(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            task_dir = Path(raw)
            source = task_dir / "source_page.png"
            Image.new("RGB", (16, 9), "white").save(source)
            refined = task_dir / "page_refined_editable.pptx"
            make_page(refined, "validated")
            baseline = task_dir / "working" / "baseline_v1"
            baseline.mkdir(parents=True)
            make_page(baseline / "recomposed_from_elements.pptx", "baseline")
            write_json(baseline / "visual_elements_manifest.json", {"elements": []})
            terminal = task_dir / "working" / "terminal_validation" / "layout_quality_report.json"
            make_terminal_report(terminal, source)
            task_json = task_dir / "task.json"
            result_json = task_dir / "worker_result.json"
            write_json(
                task_json,
                {
                    "uid": "page_001",
                    "task_dir": str(task_dir),
                    "source_image": str(source),
                    "worker_result": str(result_json),
                    "refined_pptx": str(refined),
                    "baseline_dir": str(baseline),
                    "page_terminal_validation_report": str(terminal),
                },
            )
            MODULE.command_mark_start(argparse.Namespace(task_json=task_json, restart=False))
            MODULE.command_mark_baseline(argparse.Namespace(task_json=task_json, inspected=True))
            args = argparse.Namespace(
                task_json=task_json,
                status="completed",
                pptx=None,
                terminal_validation_report=None,
                iteration_count=2,
                visual_inspection_confirmed=True,
                brand_masthead_status="not_applicable",
                figure_text_status="no_external_panel_markers",
                card_decomposition_status="no_structured_cards",
                scientific_chart_status="no_data_charts",
                corrections_applied=["tightened icon crop"],
                started_at=None,
                finished_at=None,
                notes="",
            )
            self.assertEqual(0, MODULE.command_record(args))
            result = json.loads(result_json.read_text(encoding="utf-8"))
            self.assertEqual("completed", result["status"])
            self.assertEqual("passed", result["terminal_validation_status"])
            self.assertEqual(2, result["iteration_count"])
            self.assertEqual("not_applicable", result["brand_masthead_status"])
            self.assertEqual("no_external_panel_markers", result["figure_text_status"])
            self.assertEqual("no_structured_cards", result["card_decomposition_status"])
            self.assertEqual("no_data_charts", result["scientific_chart_status"])
            self.assertNotIn("preview", result)
            self.assertNotIn("review_report", result)

    def test_record_rejects_missing_brand_and_figure_text_status(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            task_dir = Path(raw)
            source = task_dir / "source_page.png"
            Image.new("RGB", (16, 9), "white").save(source)
            refined = task_dir / "page_refined_editable.pptx"
            make_page(refined, "validated")
            baseline = task_dir / "working" / "baseline_v1"
            baseline.mkdir(parents=True)
            make_page(baseline / "recomposed_from_elements.pptx", "baseline")
            write_json(baseline / "visual_elements_manifest.json", {"elements": []})
            terminal = task_dir / "working" / "terminal_validation" / "layout_quality_report.json"
            make_terminal_report(terminal, source)
            task_json = task_dir / "task.json"
            result_json = task_dir / "worker_result.json"
            write_json(
                task_json,
                {
                    "uid": "page_001",
                    "task_dir": str(task_dir),
                    "source_image": str(source),
                    "worker_result": str(result_json),
                    "refined_pptx": str(refined),
                    "baseline_dir": str(baseline),
                    "page_terminal_validation_report": str(terminal),
                },
            )
            MODULE.command_mark_start(argparse.Namespace(task_json=task_json, restart=False))
            MODULE.command_mark_baseline(argparse.Namespace(task_json=task_json, inspected=True))
            args = argparse.Namespace(
                task_json=task_json,
                status="completed",
                pptx=None,
                terminal_validation_report=None,
                iteration_count=1,
                visual_inspection_confirmed=True,
                brand_masthead_status=None,
                figure_text_status=None,
                card_decomposition_status=None,
                scientific_chart_status=None,
                corrections_applied=["checked"],
                started_at=None,
                finished_at=None,
                notes="",
            )
            self.assertEqual(2, MODULE.command_record(args))
            result = json.loads(result_json.read_text(encoding="utf-8"))
            self.assertEqual("failed", result["status"])
            self.assertIn("brand_masthead_status_missing_or_invalid", result["issues"])
            self.assertIn("figure_text_status_missing_or_invalid", result["issues"])
            self.assertIn("card_decomposition_status_missing_or_invalid", result["issues"])
            self.assertIn("scientific_chart_status_missing_or_invalid", result["issues"])

    def test_simple_merge_preserves_order_and_geometry_without_scaling(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            first = root / "first.pptx"
            second = root / "second.pptx"
            expected_geometry = make_page(first, "first")
            make_page(second, "second", width_delta=1)
            output = root / "merged.pptx"
            MODULE.merge_refined_pages_openxml([first, second], output)
            merged = Presentation(output)
            self.assertEqual(2, len(merged.slides))
            self.assertEqual(["first", "second"], [slide.shapes[0].text for slide in merged.slides])
            actual_geometry = tuple(
                int(value)
                for value in (
                    merged.slides[1].shapes[0].left,
                    merged.slides[1].shapes[0].top,
                    merged.slides[1].shapes[0].width,
                    merged.slides[1].shapes[0].height,
                )
            )
            self.assertEqual(expected_geometry, actual_geometry)

    def test_full_slide_source_picture_is_hard_failure(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / "source.png"
            Image.new("RGB", (160, 90), "white").save(source)
            deck = root / "bad_full_slide.pptx"
            make_picture_overlap_page(deck, source, full_slide=True)
            issues, report = MODULE.refined_raster_integrity_issues(deck, source)
            self.assertTrue(any(item.startswith("full_slide_picture_forbidden") for item in issues))
            self.assertTrue(any(item.startswith("source_image_embedded_as_picture") for item in issues))
            self.assertEqual("failed", report["status"])

    def test_raster_text_residue_under_editable_text_is_hard_failure(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / "source.png"
            asset = root / "asset.png"
            Image.new("RGB", (160, 90), "white").save(source)
            Image.new("RGB", (500, 400), "lightgray").save(asset)
            deck = root / "bad_residue.pptx"
            make_picture_overlap_page(deck, asset, full_slide=False)
            with patch.object(
                MODULE,
                "detect_residual_tokens",
                return_value=[{"text": "ResidualText", "confidence": 99.0}],
            ):
                issues, report = MODULE.refined_raster_integrity_issues(deck, source)
            self.assertTrue(
                any(item.startswith("raster_text_residue_under_editable_text") for item in issues)
            )
            self.assertEqual("failed", report["status"])

    def test_internal_text_inside_large_figure_picture_is_not_a_hard_failure(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "figure.png"
            Image.new("RGB", (500, 300), "white").save(asset)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(asset), Inches(1), Inches(2), Inches(6), Inches(3))
            path = root / "figure_text.pptx"
            deck.save(path)
            with patch.object(
                MODULE,
                "detect_residual_tokens",
                return_value=[{"text": "ReadableText", "confidence": 99.0}],
            ):
                issues, report = MODULE.readable_figure_text_issues(path)
            self.assertEqual([], issues)
            self.assertEqual("passed", report["status"])
            self.assertEqual(
                "external_subfigure_markers_only",
                report["audit_scope"],
            )

    def test_structured_portrait_card_left_as_one_raster_is_hard_failure(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "card.png"
            Image.new("RGB", (330, 474), "white").save(asset)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(asset), Inches(0.5), Inches(0.5), Inches(2.2), Inches(3.16))
            path = root / "under_split_card.pptx"
            deck.save(path)
            tokens = [
                {"text": "Low water binder", "confidence": 99.0, "left": 20, "top": 12, "width": 170, "height": 24},
                {"text": "lower sorptivity", "confidence": 99.0, "left": 85, "top": 344, "width": 150, "height": 20},
                {"text": "higher tensile resistance", "confidence": 99.0, "left": 70, "top": 414, "width": 210, "height": 20},
            ]
            with patch.object(MODULE, "detect_residual_tokens", return_value=tokens):
                issues, report = MODULE.structured_card_decomposition_issues(path)
            self.assertTrue(
                any(item.startswith("structured_information_card_not_decomposed") for item in issues)
            )
            self.assertEqual("failed", report["status"])

    def test_portrait_scientific_panel_with_one_bottom_caption_is_not_card_failure(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "heatmap.png"
            Image.new("RGB", (330, 474), "white").save(asset)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(asset), Inches(0.5), Inches(0.5), Inches(2.2), Inches(3.16))
            path = root / "scientific_panel.pptx"
            deck.save(path)
            tokens = [
                {"text": "panel b", "confidence": 99.0, "left": 10, "top": 8, "width": 50, "height": 20},
                {"text": "11 A NEP", "confidence": 99.0, "left": 30, "top": 420, "width": 100, "height": 20},
            ]
            with patch.object(MODULE, "detect_residual_tokens", return_value=tokens):
                issues, report = MODULE.structured_card_decomposition_issues(path)
            self.assertEqual([], issues)
            self.assertEqual("passed", report["status"])

    def test_tagged_complete_scientific_chart_with_safe_margin_passes_and_skips_card_qa(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "safe_chart.png"
            image = Image.new("RGB", (500, 360), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((28, 24, 472, 330), outline="black", width=2)
            draw.line((48, 300, 450, 80), fill="blue", width=3)
            image.save(asset)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            picture = slide.shapes.add_picture(str(asset), Inches(1), Inches(1), Inches(5), Inches(3.6))
            picture.name = "SCI_CHART_COMPLETE::panel_a"
            path = root / "safe_chart.pptx"
            deck.save(path)
            chart_issues, chart_report = MODULE.scientific_chart_completeness_issues(path)
            self.assertEqual([], chart_issues)
            self.assertEqual("passed", chart_report["status"])
            with patch.object(MODULE, "detect_residual_tokens", return_value=[]):
                card_issues, card_report = MODULE.structured_card_decomposition_issues(path)
            self.assertEqual([], card_issues)
            self.assertEqual("explicit_complete_scientific_visual_tag", card_report["checks"][0]["skipped_reason"])

    def test_chart_source_box_rejects_padding_or_altered_crop(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / "source.png"
            source_image = Image.new("RGB", (640, 480), "white")
            source_draw = ImageDraw.Draw(source_image)
            source_draw.rectangle((110, 90, 529, 389), outline="black", width=2)
            source_draw.text((130, 100), "Chart title", fill="black")
            source_draw.line((140, 350, 500, 130), fill="blue", width=3)
            source_image.save(source)
            exact = source_image.crop((100, 80, 540, 400))
            exact_path = root / "exact.png"
            exact.save(exact_path)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            picture = slide.shapes.add_picture(str(exact_path), Inches(1), Inches(1), Inches(5), Inches(3.6))
            picture.name = "SCI_CHART_COMPLETE::panel_a::SRCBOX=100,80,440,320"
            exact_deck = root / "exact.pptx"
            deck.save(exact_deck)
            issues, report = MODULE.scientific_chart_completeness_issues(
                exact_deck, source, require_source_provenance=True
            )
            self.assertEqual([], issues)
            self.assertEqual("passed", report["status"])

            padded_path = root / "padded.png"
            padded = Image.new("RGB", (480, 360), "white")
            padded.paste(exact, (20, 20))
            padded.save(padded_path)
            padded_deck = Presentation()
            padded_slide = padded_deck.slides.add_slide(padded_deck.slide_layouts[6])
            padded_picture = padded_slide.shapes.add_picture(
                str(padded_path), Inches(1), Inches(1), Inches(5), Inches(3.6)
            )
            padded_picture.name = "SCI_CHART_COMPLETE::panel_a::SRCBOX=100,80,440,320"
            padded_pptx = root / "padded.pptx"
            padded_deck.save(padded_pptx)
            issues, report = MODULE.scientific_chart_completeness_issues(
                padded_pptx, source, require_source_provenance=True
            )
            self.assertTrue(
                any(item.startswith("scientific_chart_source_crop_dimensions_mismatch") for item in issues)
            )
            self.assertEqual("failed", report["status"])

    def test_panel_marker_text_must_match_tag(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            marker = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.4))
            marker.name = "SCI_PANEL_MARKER::panel_c::(c)"
            marker.text = "(EUR)"
            path = root / "bad_marker.pptx"
            deck.save(path)
            issues, report = MODULE.scientific_panel_marker_issues(path)
            self.assertTrue(
                any(item.startswith("scientific_panel_marker_text_mismatch") for item in issues)
            )
            self.assertEqual("failed", report["status"])

    def test_openxml_merge_preserves_picture_names(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "chart.png"
            Image.new("RGB", (120, 80), "white").save(asset)
            page = root / "page.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            picture = slide.shapes.add_picture(str(asset), Inches(1), Inches(1), Inches(3), Inches(2))
            picture.name = "SCI_CHART_COMPLETE::panel_a::SRCBOX=1,2,120,80"
            deck.save(page)
            output = root / "merged.pptx"
            MODULE.merge_refined_pages_openxml([page, page], output)
            merged = Presentation(output)
            names = [
                str(shape.name or "")
                for merged_slide in merged.slides
                for shape in merged_slide.shapes
                if shape.shape_type == MODULE.MSO_SHAPE_TYPE.PICTURE
            ]
            self.assertEqual(2, names.count("SCI_CHART_COMPLETE::panel_a::SRCBOX=1,2,120,80"))

    def test_tagged_scientific_chart_touching_edge_or_fragmented_fails(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            asset = root / "clipped_chart.png"
            image = Image.new("RGB", (500, 360), "white")
            draw = ImageDraw.Draw(image)
            draw.line((0, 0, 499, 0), fill="black", width=8)
            draw.line((0, 0, 0, 359), fill="black", width=8)
            image.save(asset)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            picture = slide.shapes.add_picture(str(asset), Inches(1), Inches(1), Inches(5), Inches(3.6))
            picture.name = "SCI_CHART_COMPLETE::panel_b"
            fragment = slide.shapes.add_picture(str(asset), Inches(7), Inches(1), Inches(2), Inches(1.4))
            fragment.name = "SCI_CHART_FRAGMENT::panel_b"
            path = root / "clipped_chart.pptx"
            deck.save(path)
            issues, report = MODULE.scientific_chart_completeness_issues(path)
            self.assertTrue(any(item.startswith("scientific_chart_missing_safe_margin") for item in issues))
            self.assertTrue(any(item.startswith("scientific_chart_fragment_forbidden") for item in issues))
            self.assertEqual("failed", report["status"])

    def test_three_page_prepare_record_watch_merge_and_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            inputs_dir = root / "inputs"
            inputs_dir.mkdir()
            inputs = []
            for index in range(1, 4):
                source = inputs_dir / f"slide_{index:02d}.png"
                Image.new("RGB", (160, 90), (240, 240 - index, 255)).save(source)
                inputs.append(source)
            outdir = root / "output"
            preflight_report = root / "output.preflight.json"
            preflight_args = argparse.Namespace(
                inputs=inputs,
                report=preflight_report,
                intended_outdir=outdir,
                available_worker_slots=3,
                root_task_count=1,
                cpu_threads_per_worker=4,
                memory_gb_per_worker=4.0,
                powerpoint_worker_cap=4,
            )
            with patch.object(
                MODULE,
                "detect_system_capacity",
                return_value={
                    "assessment_kind": "test",
                    "logical_cpu_count": 16,
                    "total_memory_gb": 32.0,
                    "available_memory_gb": 24.0,
                    "active_powerpoint_processes": 0,
                    "assumptions": {},
                    "capacity_by_resource": {"cpu": 4, "available_memory": 6, "powerpoint": 4},
                    "hardware_safe_worker_limit": 4,
                },
            ):
                self.assertEqual(0, MODULE.command_preflight(preflight_args))
            preflight = json.loads(preflight_report.read_text(encoding="utf-8"))
            self.assertEqual("awaiting_user_confirmation", preflight["status"])
            self.assertEqual(3, preflight["summary"]["full_slide_raster_page_count"])
            self.assertEqual([1, 2, 3], preflight["parallelism_confirmation"]["allowed_workers"])
            self.assertEqual(3, preflight["parallelism_confirmation"]["recommended_workers"])
            self.assertEqual(4, preflight["parallelism_confirmation"]["hardware_safe_worker_limit"])
            self.assertEqual(3, preflight["parallelism_confirmation"]["maximum_allowed_workers"])
            prepare_args = argparse.Namespace(
                inputs=inputs,
                outdir=outdir,
                workers=3,
                python=Path(sys.executable),
                v1_skill=MODULE.DEFAULT_V1_SKILL,
                worker_contract=MODULE.DEFAULT_WORKER_CONTRACT,
                extractor=MODULE.DEFAULT_EXTRACTOR,
                validator=MODULE.DEFAULT_VALIDATOR,
                dpi=300,
                timeout=60,
                preflight_report=preflight_report,
                confirmed=False,
                capacity_override_confirmed=False,
                clean=True,
            )
            with self.assertRaisesRegex(RuntimeError, "not authorized"):
                MODULE.command_prepare(prepare_args)
            self.assertFalse(outdir.exists())
            prepare_args.confirmed = True
            prepare_args.workers = 4
            with self.assertRaisesRegex(RuntimeError, "not confirmed by preflight"):
                MODULE.command_prepare(prepare_args)
            self.assertFalse(outdir.exists())
            prepare_args.workers = 3
            self.assertEqual(0, MODULE.command_prepare(prepare_args))
            manifest_path = outdir / "parallel_v1_manifest.json"
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(3, manifest["page_count"])
            consumed = json.loads(preflight_report.read_text(encoding="utf-8"))
            self.assertEqual("confirmed_and_prepared", consumed["status"])
            self.assertEqual(3, consumed["parallelism_confirmation"]["selected_workers"])
            self.assertFalse((outdir / "parallel_v1_manifest.csv").exists())

            for index, task in enumerate(manifest["tasks"], start=1):
                request_text = Path(task["worker_request"]).read_text(encoding="utf-8")
                self.assertIn("do not reread either skill file", request_text)
                self.assertIn("--no-zip", request_text)
                self.assertIn("SCI_CHART_COMPLETE::", request_text)
                self.assertIn("::SRCBOX=", request_text)
                self.assertIn("SCI_PANEL_MARKER::", request_text)
                self.assertIn("no transparent/white padding", request_text)
                self.assertIn("--scientific-chart-status", request_text)
                self.assertIn("Never split one continuous chart", request_text)
                self.assertIn("recreated as editable text", request_text)
                self.assertNotIn("page_review.json", request_text)
                task_json = Path(task["task_json"])
                MODULE.command_mark_start(argparse.Namespace(task_json=task_json, restart=False))
                baseline = Path(task["baseline_dir"])
                baseline.mkdir(parents=True, exist_ok=True)
                make_page(baseline / "recomposed_from_elements.pptx", "baseline")
                write_json(baseline / "visual_elements_manifest.json", {"elements": []})
                MODULE.command_mark_baseline(argparse.Namespace(task_json=task_json, inspected=True))

                result_path = Path(task["worker_result"])
                running = json.loads(result_path.read_text(encoding="utf-8"))
                running["started_at"] = f"2026-01-01T10:00:{index * 10:02d}+00:00"
                write_json(result_path, running)
                make_page(Path(task["refined_pptx"]), f"page-{index}", width_delta=index - 1)
                terminal = Path(task["page_terminal_validation_report"])
                make_terminal_report(terminal, Path(task["source_image"]))
                record_args = argparse.Namespace(
                    task_json=task_json,
                    status="completed",
                    pptx=None,
                    terminal_validation_report=None,
                    iteration_count=index,
                    visual_inspection_confirmed=True,
                    brand_masthead_status="not_applicable",
                    figure_text_status="no_external_panel_markers",
                    card_decomposition_status="no_structured_cards",
                    scientific_chart_status="no_data_charts",
                    corrections_applied=[f"page {index} checked"],
                    started_at=None,
                    finished_at=f"2026-01-01T10:02:{index * 10:02d}+00:00",
                    notes="",
                )
                self.assertEqual(0, MODULE.command_record(record_args))

            watch_args = argparse.Namespace(
                outdir=outdir,
                poll_seconds=0.1,
                timeout=2.0,
                merged_output="merged_v1_refined_editable.pptx",
                min_speedup=1.4,
                allow_timing_miss=False,
                keep_working_artifacts=False,
            )
            self.assertEqual(0, MODULE.command_watch_finalize(watch_args))
            merged = Presentation(outdir / "merged_v1_refined_editable.pptx")
            self.assertEqual(3, len(merged.slides))
            self.assertEqual(
                {"merged_v1_refined_editable.pptx", "parallel_timing_report.json", "parallel_v1_finalization.json"},
                {path.name for path in outdir.iterdir()},
            )
            final = json.loads((outdir / "parallel_v1_finalization.json").read_text(encoding="utf-8"))
            self.assertEqual("completed", final["status"])
            self.assertEqual("passed", final["merge_integrity_status"])
            self.assertEqual("passed", final["cleanup"]["status"])
            timing = json.loads((outdir / "parallel_timing_report.json").read_text(encoding="utf-8"))
            self.assertEqual(3, timing["peak_concurrent_workers"])
            self.assertGreaterEqual(timing["measured_speedup"], 1.4)

    def test_six_worker_multi_root_prepare_requires_override_and_splits_three_plus_three(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            inputs = []
            for index in range(1, 7):
                source = root / f"slide_{index:02d}.png"
                Image.new("RGB", (160, 90), (240, 240, 250 - index)).save(source)
                inputs.append(source)
            outdir = root / "output"
            preflight_report = root / "output.preflight.json"
            preflight_args = argparse.Namespace(
                inputs=inputs,
                report=preflight_report,
                intended_outdir=outdir,
                available_worker_slots=3,
                root_task_count=2,
                cpu_threads_per_worker=4,
                memory_gb_per_worker=4.0,
                powerpoint_worker_cap=4,
            )
            capacity = {
                "assessment_kind": "test",
                "logical_cpu_count": 16,
                "total_memory_gb": 32.0,
                "available_memory_gb": 16.0,
                "active_powerpoint_processes": 0,
                "assumptions": {},
                "capacity_by_resource": {"cpu": 4, "available_memory": 4, "powerpoint": 4},
                "hardware_safe_worker_limit": 3,
            }
            with patch.object(MODULE, "detect_system_capacity", return_value=capacity):
                self.assertEqual(0, MODULE.command_preflight(preflight_args))
            report = json.loads(preflight_report.read_text(encoding="utf-8"))
            parallelism = report["parallelism_confirmation"]
            self.assertEqual(3, parallelism["recommended_workers"])
            self.assertEqual(6, parallelism["maximum_allowed_workers"])
            self.assertEqual([4, 5, 6], parallelism["capacity_override_required_for"])

            prepare_args = argparse.Namespace(
                inputs=inputs,
                outdir=outdir,
                workers=6,
                python=Path(sys.executable),
                v1_skill=MODULE.DEFAULT_V1_SKILL,
                worker_contract=MODULE.DEFAULT_WORKER_CONTRACT,
                extractor=MODULE.DEFAULT_EXTRACTOR,
                validator=MODULE.DEFAULT_VALIDATOR,
                dpi=300,
                timeout=60,
                preflight_report=preflight_report,
                confirmed=True,
                capacity_override_confirmed=False,
                clean=True,
            )
            with self.assertRaisesRegex(RuntimeError, "capacity override"):
                MODULE.command_prepare(prepare_args)
            self.assertFalse(outdir.exists())
            prepare_args.capacity_override_confirmed = True
            self.assertEqual(0, MODULE.command_prepare(prepare_args))
            manifest = json.loads((outdir / "parallel_v1_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(6, manifest["worker_limit"])
            self.assertEqual(2, manifest["root_task_count"])
            self.assertEqual([3, 3], [len(shard["task_uids"]) for shard in manifest["shards"]])
            self.assertEqual(
                ["page_001_source_001_p001", "page_002_source_002_p001", "page_003_source_003_p001"],
                manifest["shards"][0]["task_uids"],
            )


if __name__ == "__main__":
    unittest.main()
