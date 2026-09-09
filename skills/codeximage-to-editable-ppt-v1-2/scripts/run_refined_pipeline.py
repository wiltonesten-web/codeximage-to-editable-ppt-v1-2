#!/usr/bin/env python3
"""Incremental end-to-end refined rebuild scheduler.

The scheduler separates lightweight source-reference extraction, page-specific
refined builds, PowerPoint geometry preflight, merge, and final visual validation. Expensive
work is content-addressed, only dirty pages rebuild, PowerPoint automation is
serialized, and full image comparison runs once on the final ordered deck.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
from copy import deepcopy
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_RUN_BATCHES = SCRIPT_DIR / "run_batches.py"
DEFAULT_VALIDATOR = SCRIPT_DIR / "validate_layout_powerpoint.py"
DEFAULT_AUTOFIT = SCRIPT_DIR / "autofit_powerpoint_text.py"
DEFAULT_ASSET_TEXT_AUDITOR = SCRIPT_DIR / "audit_png_text_residue.py"


@dataclass
class PageTask:
    uid: str
    task_index: int
    source_index: int
    source_input: str
    page_index: int
    page_id: str
    source_image: str
    reference_dir: str
    output_dir: str
    pptx: str
    refine_status: str = "pending"
    refine_cached: bool = False
    asset_text_status: str = "pending"
    geometry_status: str = "pending"
    geometry_retries: int = 0
    final_status: str = "pending"
    issues: str = ""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _stable_hash(parts: Iterable[Any]) -> str:
    payload = json.dumps(list(parts), ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _load_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _run(cmd: list[str], log: Path, timeout: int) -> int:
    log.parent.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    for name in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
        env[name] = "1"
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, env=env, timeout=timeout)
    log.write_text(proc.stdout or "", encoding="utf-8")
    return int(proc.returncode)


def _reference_key(args: argparse.Namespace, inputs: list[Path]) -> str:
    return _stable_hash([
        "source-reference-v1",
        *[(str(path.resolve()), _sha256_file(path.resolve())) for path in inputs],
        args.dpi,
        _sha256_file(args.run_batches.resolve()),
    ])


def ensure_reference(args: argparse.Namespace, inputs: list[Path], cache: dict) -> tuple[Path, bool]:
    reference = args.outdir.resolve() / "source_reference"
    manifest = reference / "page_task_manifest.json"
    key = _reference_key(args, inputs)
    if not args.force_reference and cache.get("reference_key") == key and manifest.exists():
        return reference, True
    cmd = [
        str(args.python.resolve()), str(args.run_batches.resolve()),
        *[str(path.resolve()) for path in inputs],
        "--outdir", str(reference),
        "--batch-size", "1",
        "--batch-workers", str(args.reference_workers),
        "--workers", str(args.reference_workers),
        "--python", str(args.python.resolve()),
        "--dpi", str(args.dpi),
        "--dry-run", "--clean-batch-root",
    ]
    rc = _run(cmd, args.outdir.resolve() / "logs" / "source_reference.log", args.reference_timeout)
    if rc != 0 or not manifest.exists():
        raise RuntimeError(f"Source-reference extraction failed; see {args.outdir.resolve() / 'logs' / 'source_reference.log'}")
    cache["reference_key"] = key
    cache["pages"] = {}
    return reference, False


def load_tasks(args: argparse.Namespace, reference: Path) -> list[PageTask]:
    rows = _load_json(reference / "page_task_manifest.json", [])
    tasks: list[PageTask] = []
    for row in rows:
        source_index = int(row["source_index"])
        page_index = int(row["page_index"])
        uid = f"source_{source_index:03d}_page_{page_index:03d}"
        output = args.outdir.resolve() / "refined_pages" / uid
        tasks.append(PageTask(
            uid=uid,
            task_index=int(row["task_index"]),
            source_index=source_index,
            source_input=str(row["source_input"]),
            page_index=page_index,
            page_id=str(row["page_id"]),
            source_image=str(Path(row["page_image_path"]).resolve()),
            reference_dir=str(Path(row["extraction_dir"]).resolve()),
            output_dir=str(output),
            pptx=str(output / "refined.pptx"),
        ))
    return sorted(tasks, key=lambda item: (item.source_index, item.page_index))


def _refine_key(args: argparse.Namespace, task: PageTask) -> str:
    source = Path(task.source_image)
    script_hash = "automatic_reference_only"
    if args.refine_script:
        script_hash = _sha256_file(args.refine_script.resolve())
    return _stable_hash([
        "refine-v3", task.uid, _sha256_file(source),
        script_hash, args.refine_extra_arg, _refine_dependency_hashes(args, task),
        args.asset_text_audit,
        _sha256_file(args.asset_text_auditor.resolve()) if args.asset_text_audit else "",
        args.asset_text_min_confidence,
        args.ocr_lang,
    ])


def _refine_dependency_hashes(args: argparse.Namespace, task: PageTask) -> list[tuple[str, str]]:
    """Hash page-specific files referenced by builder arguments, especially JSON specs."""
    hashes: list[tuple[str, str]] = []
    values = list(args.refine_extra_arg)
    for index, value in enumerate(values):
        if value == "--spec-dir" and index + 1 < len(values):
            spec = Path(values[index + 1]) / f"{Path(task.source_input).stem}.json"
            if spec.is_file():
                hashes.append((str(spec.resolve()), _sha256_file(spec.resolve())))
        else:
            candidate = Path(value)
            if candidate.is_file():
                hashes.append((str(candidate.resolve()), _sha256_file(candidate.resolve())))
    return hashes


def _refine_one(args: argparse.Namespace, task: PageTask, cache_pages: dict) -> tuple[PageTask, str]:
    output = Path(task.output_dir)
    pptx = Path(task.pptx)
    key = _refine_key(args, task)
    old = cache_pages.get(task.uid, {})
    if not args.force_refine and old.get("refine_key") == key and pptx.exists():
        task.refine_status = "completed"
        task.refine_cached = True
        audit_payload = _load_json(output / "asset_text_residue_report.json", {})
        task.asset_text_status = str(audit_payload.get("status", "not_run"))
        return task, key
    partial = output.with_name(f"{output.name}.partial")
    failed = output.with_name(f"{output.name}.failed")
    if partial.exists():
        shutil.rmtree(partial)
    partial.mkdir(parents=True, exist_ok=True)
    target = partial / "refined.pptx"
    if args.automatic_output_only:
        build_source_reference_pptx(Path(task.source_image), target)
        (partial / "refine.log").write_text("automatic_output_only: built exact source-reference PPTX\n", encoding="utf-8")
        rc = 0
    else:
        cmd = [
            str(args.python.resolve()), str(args.refine_script.resolve()),
            "--source", task.source_image,
            "--reference-dir", task.reference_dir,
            "--outdir", str(partial),
            "--pptx-out", str(target),
            "--page-id", task.page_id,
            "--page-index", str(task.page_index),
            "--source-input", task.source_input,
            *args.refine_extra_arg,
        ]
        rc = _run(cmd, partial / "refine.log", args.refine_timeout)
    failure_reason = "refine_failed"
    if rc == 0 and target.exists() and args.asset_text_audit and not args.automatic_output_only:
        audit_report = partial / "asset_text_residue_report.json"
        audit_cmd = [
            str(args.python.resolve()),
            str(args.asset_text_auditor.resolve()),
            "--root", str(partial),
            "--report", str(audit_report),
            "--lang", args.ocr_lang,
            "--min-confidence", str(args.asset_text_min_confidence),
        ]
        allowlist = partial / "asset_text_allowlist.json"
        if allowlist.exists():
            audit_cmd.extend(["--allowlist", str(allowlist)])
        audit_rc = _run(audit_cmd, partial / "asset_text_residue_audit.log", args.refine_timeout)
        audit_payload = _load_json(audit_report, {})
        task.asset_text_status = str(audit_payload.get("status", "auditor_error"))
        if audit_rc != 0:
            rc = audit_rc
            failure_reason = "asset_text_residue_detected"
    elif rc == 0 and target.exists():
        task.asset_text_status = "not_run"
    if rc != 0 or not target.exists():
        if failed.exists():
            shutil.rmtree(failed)
        partial.replace(failed)
        task.refine_status = "failed"
        task.issues = failure_reason
        return task, key
    if output.exists():
        shutil.rmtree(output)
    partial.replace(output)
    task.refine_status = "completed"
    return task, key


def refine_pages(args: argparse.Namespace, tasks: list[PageTask], cache: dict) -> None:
    pages = cache.setdefault("pages", {})
    with concurrent.futures.ThreadPoolExecutor(max_workers=args.refine_workers) as executor:
        futures = {executor.submit(_refine_one, args, task, pages): task for task in tasks}
        for future in concurrent.futures.as_completed(futures):
            task, key = future.result()
            pages.setdefault(task.uid, {})["refine_key"] = key
            pages[task.uid]["refine_status"] = task.refine_status
    if any(task.refine_status != "completed" for task in tasks):
        raise RuntimeError("One or more refined page builds failed")


def build_source_reference_pptx(source: Path, target: Path) -> None:
    """Create an exact one-picture reference page; never present it as editable final output."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    with Image.open(source) as image:
        width_px, height_px = image.size
    slide_width = 13.333333
    source_ratio = width_px / height_px
    slide_height = 7.5 if abs(source_ratio - (16 / 9)) <= 0.02 else slide_width / source_ratio
    prs = Presentation()
    prs.slide_width = Inches(slide_width)
    prs.slide_height = Inches(slide_height)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(source.resolve()), 0, 0, width=prs.slide_width, height=prs.slide_height)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(target)


def _geometry_key(args: argparse.Namespace, pptx: Path) -> str:
    return _stable_hash(["geometry-v2", _sha256_file(pptx), args.text_tolerance, _sha256_file(args.validator.resolve())])


def geometry_preflight(args: argparse.Namespace, tasks: list[PageTask], cache: dict) -> None:
    pages = cache.setdefault("pages", {})
    for task in tasks:  # PowerPoint COM is intentionally serialized.
        pptx = Path(task.pptx)
        report_dir = Path(task.output_dir) / "geometry_validation"
        key = _geometry_key(args, pptx)
        old = pages.setdefault(task.uid, {})
        report = report_dir / "layout_quality_report.json"
        if not args.force_validate and old.get("geometry_key") == key and report.exists():
            payload = _load_json(report, {})
            task.geometry_status = "passed" if payload.get("status") == "passed" else "manual_review_required"
            continue
        for attempt in range(args.max_geometry_retries + 1):
            cmd = [
                str(args.python.resolve()), str(args.validator.resolve()), str(pptx),
                "--outdir", str(report_dir), "--geometry-only",
                "--text-tolerance", str(args.text_tolerance),
            ]
            rc = _run(cmd, report_dir / f"geometry_attempt_{attempt + 1}.log", args.powerpoint_timeout)
            payload = _load_json(report, {})
            slides = payload.get("slides", [])
            overflow = sum(int(row.get("text_overflow_count", 0)) for row in slides)
            off_slide = sum(int(row.get("off_slide_shape_count", 0)) for row in slides)
            task.geometry_retries = attempt
            if rc == 0:
                task.geometry_status = "passed"
                break
            if not args.autofit_text or overflow == 0 or off_slide > 0 or attempt >= args.max_geometry_retries:
                task.geometry_status = "manual_review_required"
                task.issues = ";".join(filter(None, [task.issues, "geometry_failed"]))
                break
            fix_cmd = [
                str(args.python.resolve()), str(args.autofit_script.resolve()), str(pptx), "--inplace",
                "--text-tolerance", str(args.text_tolerance),
                "--report", str(report_dir / f"autofit_attempt_{attempt + 1}.json"),
            ]
            fix_rc = _run(fix_cmd, report_dir / f"autofit_attempt_{attempt + 1}.log", args.powerpoint_timeout)
            if fix_rc != 0:
                task.geometry_status = "manual_review_required"
                task.issues = ";".join(filter(None, [task.issues, "autofit_failed"]))
                break
        old["geometry_key"] = _geometry_key(args, pptx)
        old["geometry_status"] = task.geometry_status


def merge_page_decks_openxml(page_files: list[Path], output: Path) -> None:
    if not page_files:
        raise RuntimeError("No refined page PPTX files to merge")
    output.parent.mkdir(parents=True, exist_ok=True)
    if len(page_files) == 1:
        shutil.copy2(page_files[0], output)
        return
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    first = Presentation(page_files[0])
    merged = Presentation()
    merged.slide_width = first.slide_width
    merged.slide_height = first.slide_height
    blank = merged.slide_layouts[6]
    for page_file in page_files:
        source_deck = Presentation(page_file)
        if len(source_deck.slides) != 1:
            raise RuntimeError(f"Expected one slide in refined page deck: {page_file}")
        if source_deck.slide_width != merged.slide_width or source_deck.slide_height != merged.slide_height:
            raise RuntimeError(f"Slide size mismatch in refined page deck: {page_file}")
        source_slide = source_deck.slides[0]
        target_slide = merged.slides.add_slide(blank)
        source_bg = source_slide._element.cSld.bg
        if source_bg is not None:
            target_bg = target_slide._element.cSld.bg
            if target_bg is not None:
                target_slide._element.cSld.remove(target_bg)
            target_slide._element.cSld.insert(0, deepcopy(source_bg))
        for shape in source_slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture = target_slide.shapes.add_picture(
                    BytesIO(shape.image.blob), shape.left, shape.top, shape.width, shape.height
                )
                picture.crop_left = shape.crop_left
                picture.crop_right = shape.crop_right
                picture.crop_top = shape.crop_top
                picture.crop_bottom = shape.crop_bottom
                picture.rotation = shape.rotation
            else:
                target_slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    merged.save(output)


def merge_if_dirty(args: argparse.Namespace, tasks: list[PageTask], cache: dict) -> tuple[Path, bool]:
    page_files = [Path(task.pptx) for task in tasks]
    output = args.outdir.resolve() / args.merged_output
    key = _stable_hash(["merge-v3-openxml", *[_sha256_file(path) for path in page_files]])
    if not args.force_merge and cache.get("merge_key") == key and output.exists():
        return output, True
    merge_page_decks_openxml(page_files, output)
    cache["merge_key"] = key
    cache.pop("final_validation_key", None)
    return output, False


def final_validation(args: argparse.Namespace, tasks: list[PageTask], merged: Path, cache: dict) -> tuple[dict, bool]:
    report_dir = args.outdir.resolve() / "final_powerpoint_validation"
    report = report_dir / "layout_quality_report.json"
    key = _stable_hash([
        "final-v2", _sha256_file(merged),
        *[_sha256_file(Path(task.source_image)) for task in tasks],
        args.max_mae, args.max_changed_ratio, args.diff_threshold, args.text_tolerance,
        _sha256_file(args.validator.resolve()),
    ])
    if not args.force_validate and cache.get("final_validation_key") == key and report.exists():
        return _load_json(report, {}), True
    cmd = [
        str(args.python.resolve()), str(args.validator.resolve()), str(merged),
        "--outdir", str(report_dir),
        "--max-mae", str(args.max_mae),
        "--max-changed-ratio", str(args.max_changed_ratio),
        "--diff-threshold", str(args.diff_threshold),
        "--text-tolerance", str(args.text_tolerance),
    ]
    for task in tasks:
        cmd.extend(["--source", task.source_image])
    _run(cmd, report_dir / "final_validation.log", args.powerpoint_timeout)
    payload = _load_json(report, {"status": "validator_error"})
    cache["final_validation_key"] = key
    return payload, False


def write_summary(args: argparse.Namespace, tasks: list[PageTask], payload: dict, reuse: dict) -> None:
    slides = payload.get("slides", [])
    for index, task in enumerate(tasks):
        if index < len(slides):
            task.final_status = str(slides[index].get("status", "manual_review_required"))
            final_issues = str(slides[index].get("issues", ""))
            task.issues = ";".join(filter(None, [task.issues, final_issues]))
    summary = {
        "status": payload.get("status", "failed_or_review_required"),
        "reuse": reuse,
        "policy": {
            "reference_parallelism": args.reference_workers,
            "refine_parallelism": args.refine_workers,
            "powerpoint_parallelism": 1,
            "max_geometry_retries": args.max_geometry_retries,
            "final_visual_validation_runs": 1,
        },
        "pages": [asdict(task) for task in tasks],
        "final_validation_report": str(args.outdir.resolve() / "final_powerpoint_validation" / "layout_quality_report.json"),
    }
    _write_json(args.outdir.resolve() / "refined_pipeline_summary.json", summary)
    rows = [asdict(task) for task in tasks]
    with (args.outdir.resolve() / "refined_pipeline_summary.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader(); writer.writerows(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Incremental refined image-to-editable-PPT pipeline.")
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--refine-script", type=Path, help="Page builder implementing the refined page CLI contract.")
    parser.add_argument("--automatic-output-only", action="store_true", help="Build exact source-image reference pages; not an editable final deck.")
    parser.add_argument("--refine-extra-arg", action="append", default=[])
    parser.add_argument("--python", type=Path, default=Path(sys.executable))
    parser.add_argument("--run-batches", type=Path, default=DEFAULT_RUN_BATCHES)
    parser.add_argument("--validator", type=Path, default=DEFAULT_VALIDATOR)
    parser.add_argument("--autofit-script", type=Path, default=DEFAULT_AUTOFIT)
    parser.add_argument("--asset-text-auditor", type=Path, default=DEFAULT_ASSET_TEXT_AUDITOR)
    parser.add_argument("--reference-workers", "--baseline-workers", dest="reference_workers", type=int, default=4, choices=range(1, 5))
    parser.add_argument("--refine-workers", type=int, default=4, choices=range(1, 5))
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--granularity", choices=("coarse", "normal", "fine", "ultra"), default="fine")
    parser.add_argument("--ocr", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--ocr-lang", default="chi_sim+eng")
    parser.add_argument("--ocr-confidence-threshold", type=float, default=75.0)
    parser.add_argument("--asset-text-audit", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--asset-text-min-confidence", type=float, default=60.0)
    parser.add_argument("--default-font-family", default="Microsoft YaHei")
    parser.add_argument("--autofit-text", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--max-geometry-retries", type=int, default=2, choices=range(0, 3))
    parser.add_argument("--text-tolerance", type=float, default=1.03)
    parser.add_argument("--max-mae", type=float, default=15.0)
    parser.add_argument("--max-changed-ratio", type=float, default=0.15)
    parser.add_argument("--diff-threshold", type=int, default=25)
    parser.add_argument("--merged-output", default="merged_refined_editable.pptx")
    parser.add_argument("--force-reference", "--force-baseline", dest="force_reference", action="store_true")
    parser.add_argument("--force-refine", action="store_true")
    parser.add_argument("--force-merge", action="store_true")
    parser.add_argument("--force-validate", action="store_true")
    parser.add_argument("--reference-timeout", "--baseline-timeout", dest="reference_timeout", type=int, default=1800)
    parser.add_argument("--refine-timeout", type=int, default=1800)
    parser.add_argument("--powerpoint-timeout", type=int, default=300)
    args = parser.parse_args()
    if not args.automatic_output_only and args.refine_script is None:
        parser.error("--refine-script is required unless --automatic-output-only is explicitly selected")
    if args.automatic_output_only and args.merged_output == "merged_refined_editable.pptx":
        args.merged_output = "merged_source_reference.pptx"
    return args


def main() -> int:
    args = parse_args()
    args.outdir = args.outdir.resolve(); args.outdir.mkdir(parents=True, exist_ok=True)
    inputs = [path.resolve() for path in args.inputs]
    for path in inputs:
        if not path.exists():
            raise SystemExit(f"Input does not exist: {path}")
    cache_path = args.outdir / ".refined_pipeline_cache.json"
    cache = _load_json(cache_path, {})
    try:
        reference, reference_cached = ensure_reference(args, inputs, cache)
        tasks = load_tasks(args, reference)
        refine_pages(args, tasks, cache)
        geometry_preflight(args, tasks, cache)
        merged, merge_cached = merge_if_dirty(args, tasks, cache)
        payload, validation_cached = final_validation(args, tasks, merged, cache)
        reuse = {
            "source_reference_cached": reference_cached,
            "refined_pages_cached": sum(1 for task in tasks if task.refine_cached),
            "merged_cached": merge_cached,
            "final_validation_cached": validation_cached,
        }
        write_summary(args, tasks, payload, reuse)
        _write_json(cache_path, cache)
        stale_error = args.outdir / "refined_pipeline_error.json"
        if stale_error.exists():
            stale_error.unlink()
        print(json.dumps({"merged": str(merged), "status": payload.get("status"), "reuse": reuse}, ensure_ascii=False, indent=2))
        return 0 if payload.get("status") == "passed" else 2
    except Exception as exc:
        _write_json(cache_path, cache)
        failure = {"status": "pipeline_error", "error": str(exc)}
        _write_json(args.outdir / "refined_pipeline_error.json", failure)
        print(json.dumps(failure, ensure_ascii=False, indent=2))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
